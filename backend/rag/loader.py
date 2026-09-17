import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader


HEADING_PATTERN = re.compile(r"^(?:[一二三四五六七八九十]+、|\d+(?:\.\d+)+\s*)")
PAGE_NUMBER_PATTERN = re.compile(r"^第\s*\d+\s*页$")


def _clean_pdf_lines(text: str, page_number: int) -> list[str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    # ReportLab 等 PDF 常把页眉和页码提取到正文最前面。
    if len(lines) >= 2 and PAGE_NUMBER_PATTERN.match(lines[1]):
        lines = lines[2:]
    elif lines and PAGE_NUMBER_PATTERN.match(lines[0]):
        lines = lines[1:]

    return [line for line in lines if line != f"第 {page_number} 页"]


def _lines_to_sections(lines: list[str], page_number: int | None):
    sections = []
    page_title = ""
    current_title = ""
    current_lines: list[str] = []

    def flush():
        if not current_lines:
            return
        text = "\n".join(current_lines).strip()
        if text:
            sections.append(
                {
                    "text": text,
                    "page": page_number,
                    "title": current_title or page_title,
                }
            )

    for line in lines:
        if HEADING_PATTERN.match(line):
            flush()
            current_lines = []
            current_title = line
            if "、" in line and not page_title:
                page_title = line
            continue
        current_lines.append(line)

    flush()
    return sections


def load_pdf_sections(file_path: str):
    reader = PdfReader(file_path)
    sections = []

    for page_number, page in enumerate(reader.pages, start=1):
        content = page.extract_text() or ""
        lines = _clean_pdf_lines(content, page_number)
        page_sections = _lines_to_sections(lines, page_number)

        if page_sections:
            sections.extend(page_sections)
        elif lines:
            sections.append(
                {"text": "\n".join(lines), "page": page_number, "title": ""}
            )

    return sections


def load_docx_sections(file_path: str):
    document = Document(file_path)
    sections = []
    current_title = ""

    for paragraph in document.paragraphs:
        content = paragraph.text.strip()
        if not content:
            continue

        if paragraph.style.name.startswith("Heading") or HEADING_PATTERN.match(content):
            current_title = content
            continue

        sections.append({"text": content, "page": None, "title": current_title})

    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            sections.append(
                {"text": "\n".join(rows), "page": None, "title": current_title}
            )

    return sections


def load_document_sections(file_path: str):
    suffix = Path(file_path).suffix.lower()

    if suffix == ".pdf":
        return load_pdf_sections(file_path)
    if suffix == ".docx":
        return load_docx_sections(file_path)

    raise ValueError(f"不支持的文件类型: {suffix}")


def load_pdf(file_path: str) -> str:
    return "\n\n".join(section["text"] for section in load_pdf_sections(file_path))


def load_docx(file_path: str) -> str:
    return "\n\n".join(section["text"] for section in load_docx_sections(file_path))


def load_document(file_path: str) -> str:
    return "\n\n".join(
        section["text"] for section in load_document_sections(file_path)
    )
